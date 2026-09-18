"""UI tests for the accounting screens: ledger report + exports, vouchers,
chart of accounts, and the FIN-EC-03 paid-expense warning."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from core.models import UserProfile

from .models import (
    AccountHead, Office, ExpenseCategory, OfficeExpense, Voucher, VoucherAuditLog,
    VoucherLine,
)


class AccountingUiBase(TestCase):
    def setUp(self):
        self.finance = User.objects.create_user('uifin', 'f@example.com', 'pass12345')
        UserProfile.objects.create(user=self.finance, role='accounts')
        self.supervisor = User.objects.create_user('uisup', 's@example.com', 'pass12345')
        UserProfile.objects.create(user=self.supervisor, role='management')
        self.sales = User.objects.create_user('uisal', 'x@example.com', 'pass12345')
        UserProfile.objects.create(user=self.sales, role='sales')

        self.cash = AccountHead.objects.create(code='1000', name='Cash', nature='debit')
        self.capital = AccountHead.objects.create(code='3000', name='Capital', nature='credit')

    def login(self, user):
        self.client.force_login(user)

    def posted_voucher(self):
        v = Voucher.objects.create(voucher_type='JV', date=date.today(), created_by=self.finance)
        VoucherLine.objects.create(voucher=v, account_head=self.cash, debit=Decimal('100'))
        VoucherLine.objects.create(voucher=v, account_head=self.capital, credit=Decimal('100'))
        v.post(self.finance)
        return v


class LedgerReportUiTests(AccountingUiBase):
    def test_finance_can_view_ledger_report(self):
        self.posted_voucher()
        self.login(self.finance)
        resp = self.client.get(reverse('finance_ledger'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Ledger balance')

    def test_ledger_filters_by_head_and_dates(self):
        self.posted_voucher()
        self.login(self.finance)
        resp = self.client.get(reverse('finance_ledger'), {
            'head': self.cash.pk, 'date_from': '2020-01-01', 'date_to': '2099-01-01',
        })
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Cash')

    def test_sales_denied_ledger(self):
        self.login(self.sales)
        resp = self.client.get(reverse('finance_ledger'))
        self.assertEqual(resp.status_code, 302)

    def test_excel_export(self):
        self.posted_voucher()
        self.login(self.finance)
        resp = self.client.get(reverse('finance_ledger_export_excel'), {'head': self.cash.pk})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheet', resp['Content-Type'])
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(resp.content))
        self.assertGreaterEqual(wb.active.max_row, 2)

    def test_pdf_export(self):
        self.posted_voucher()
        self.login(self.finance)
        resp = self.client.get(reverse('finance_ledger_export_pdf'), {'head': self.cash.pk})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF'))


class VoucherUiTests(AccountingUiBase):
    def test_finance_can_list_and_view_vouchers(self):
        voucher = self.posted_voucher()
        self.login(self.finance)
        self.assertEqual(self.client.get(reverse('finance_vouchers')).status_code, 200)
        detail = self.client.get(reverse('finance_voucher_detail', args=[voucher.pk]))
        self.assertEqual(detail.status_code, 200)

    def test_sales_denied_vouchers(self):
        self.login(self.sales)
        self.assertEqual(self.client.get(reverse('finance_vouchers')).status_code, 302)

    def test_create_voucher_via_ui(self):
        self.login(self.finance)
        resp = self.client.post(reverse('finance_voucher_create'), {
            'voucher_type': 'JV', 'date': date.today().isoformat(), 'narration': 'opening',
            'lines-TOTAL_FORMS': '1', 'lines-INITIAL_FORMS': '0', 'lines-MIN_NUM_FORMS': '0',
            'lines-MAX_NUM_FORMS': '1000',
            'lines-0-account_head': self.cash.pk, 'lines-0-debit': '100', 'lines-0-credit': '0',
            'lines-0-narration': '',
        })
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Voucher.objects.count(), 1)
        self.assertEqual(Voucher.objects.first().lines.count(), 1)

    def test_supervisor_can_post_voucher_via_ui(self):
        voucher = Voucher.objects.create(voucher_type='JV', date=date.today(), created_by=self.finance)
        VoucherLine.objects.create(voucher=voucher, account_head=self.cash, debit=Decimal('100'))
        VoucherLine.objects.create(voucher=voucher, account_head=self.capital, credit=Decimal('100'))
        self.login(self.supervisor)
        resp = self.client.post(reverse('finance_voucher_post', args=[voucher.pk]))
        self.assertEqual(resp.status_code, 302)
        voucher.refresh_from_db()
        self.assertEqual(voucher.status, 'posted')

    def test_finance_cannot_post_voucher_via_ui(self):
        voucher = Voucher.objects.create(voucher_type='JV', date=date.today(), created_by=self.finance)
        VoucherLine.objects.create(voucher=voucher, account_head=self.cash, debit=Decimal('100'))
        VoucherLine.objects.create(voucher=voucher, account_head=self.capital, credit=Decimal('100'))
        self.login(self.finance)
        resp = self.client.post(reverse('finance_voucher_post', args=[voucher.pk]))
        self.assertEqual(resp.status_code, 302)
        voucher.refresh_from_db()
        self.assertEqual(voucher.status, 'draft')
        # The UI must not offer the post action to a non-supervisor.
        detail = self.client.get(reverse('finance_voucher_detail', args=[voucher.pk]))
        self.assertNotContains(detail, 'Post &amp; Lock')

    def test_supervisor_unlock_requires_reason(self):
        voucher = self.posted_voucher()
        self.login(self.supervisor)
        self.client.post(reverse('finance_voucher_unlock', args=[voucher.pk]), {})
        voucher.refresh_from_db()
        self.assertTrue(voucher.is_locked)

        self.client.post(reverse('finance_voucher_unlock', args=[voucher.pk]), {'reason': 'fix'})
        voucher.refresh_from_db()
        self.assertFalse(voucher.is_locked)
        self.assertTrue(VoucherAuditLog.objects.filter(voucher=voucher, event='unlocked').exists())

    def test_finance_cannot_unlock(self):
        voucher = self.posted_voucher()
        self.login(self.finance)
        resp = self.client.post(reverse('finance_voucher_unlock', args=[voucher.pk]), {'reason': 'fix'})
        self.assertEqual(resp.status_code, 302)
        voucher.refresh_from_db()
        self.assertTrue(voucher.is_locked)


class AccountHeadUiTests(AccountingUiBase):
    def test_finance_can_view_tree(self):
        AccountHead.objects.create(code='1100', name='Petty Cash', parent=self.cash)
        self.login(self.finance)
        resp = self.client.get(reverse('finance_account_heads'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Petty Cash')

    def test_create_head_via_ui(self):
        self.login(self.finance)
        resp = self.client.post(reverse('finance_account_head_create'),
                                {'code': '4000', 'name': 'Sales Revenue', 'nature': 'credit'})
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(AccountHead.objects.filter(code='4000').exists())

    def test_sales_denied_tree(self):
        self.login(self.sales)
        self.assertEqual(self.client.get(reverse('finance_account_heads')).status_code, 302)


class PaidExpenseWarningTests(AccountingUiBase):
    def test_paid_expense_edit_warns_about_posted_voucher(self):
        office = Office.objects.create(name='HO', office_type='head_office')
        category = ExpenseCategory.objects.create(name='Rent', category_type='rent')
        expense = OfficeExpense.objects.create(
            office=office, category=category, amount=Decimal('5000'), expense_date=date.today(),
            payment_method='cash', status='paid', created_by=self.finance,
        )
        expense.post_to_ledger()
        self.login(self.finance)
        resp = self.client.get(reverse('finance_office_expense_edit', args=[expense.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'will not rewrite')
        self.assertContains(resp, reverse('finance_voucher_detail', args=[
            Voucher.objects.get(reference_type='OfficeExpense', reference_id=expense.pk).pk,
        ]))